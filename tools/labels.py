"""Dataset helpers: read open_labels.xlsx, normalise names (Latin X vs Cyrillic Х), parse times,
find video files, drop byte-identical duplicates, and merge in data/annotations.yaml."""
from __future__ import annotations

import hashlib
import re
from pathlib import Path

import yaml
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
DEFAULT_LABELS = ROOT / "open_data" / "open" / "open_labels.xlsx"
DEFAULT_VIDEOS = ROOT / "open_data" / "open" / "videos"
ANNOTATIONS = ROOT / "data" / "annotations.yaml"


def norm_name(name: str) -> str:
    """'Х-14.mov' (Cyrillic) / 'X-14' (Latin) -> 'X-14'."""
    name = Path(name).stem.strip()
    return name.replace("Х", "X").replace("х", "x").upper()


def parse_time(text) -> float | None:
    if text is None or str(text).strip() == "":
        return None
    s = str(text).lower()
    m = re.search(r"(\d+)\s*мин", s)
    minutes = int(m.group(1)) if m else 0
    m = re.search(r"(\d+(?:[.,]\d+)?)\s*сек", s)
    if m:
        sec = float(m.group(1).replace(",", "."))
    else:
        nums = re.findall(r"\d+(?:[.,]\d+)?", s)
        sec = float(nums[-1].replace(",", ".")) if nums and not minutes else 0.0
    return minutes * 60 + sec


def load_dataset(labels: Path = DEFAULT_LABELS, videos: Path = DEFAULT_VIDEOS, dedupe: bool = True) -> list[dict]:
    wb = load_workbook(labels, read_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))[1:]
    files = {norm_name(p.name): p for p in Path(videos).glob("*") if p.suffix.lower() in {".mov", ".mp4"}}
    ann = load_annotations()

    items, seen = [], {}
    for name, ts, verdict in rows:
        if not name:
            continue
        key = norm_name(str(name))
        path = files.get(key)
        if path is None:
            continue
        positive = "без" not in str(verdict).lower()
        item = {"name": key, "path": path, "violation": positive, "t": parse_time(ts) if positive else None}
        item.update({k: v for k, v in ann.get(key, {}).items() if k not in ("name",)})
        digest = hashlib.sha1(path.read_bytes()).hexdigest()
        item["sha1"] = digest
        if dedupe and digest in seen:
            first = seen[digest]
            # same video labelled twice (X-4 / X-47): keep one item with all label times
            if positive and item["t"] is not None:
                first.setdefault("extra_t", []).append(item["t"])
            first.setdefault("duplicates", []).append(key)
            continue
        seen[digest] = item
        items.append(item)
    return items


def load_annotations() -> dict[str, dict]:
    if not ANNOTATIONS.exists():
        return {}
    data = yaml.safe_load(ANNOTATIONS.read_text(encoding="utf-8")) or {}
    return {norm_name(k): (v or {}) for k, v in data.get("clips", {}).items()}


if __name__ == "__main__":
    for it in load_dataset():
        print(f"{it['name']:6} {'VIOL' if it['violation'] else 'ok  '} t={it['t']} "
              f"type={it.get('type', '-')} setup={it.get('setup', '-')} dup={it.get('duplicates', '')}")
